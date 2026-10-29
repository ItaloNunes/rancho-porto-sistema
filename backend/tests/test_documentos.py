"""Contrato e recibo: preenchimento a partir do plano de pagamento e
conversão pra PDF (app/docx_pdf.py).

O teste mais importante aqui é o de equivalência de texto: o PDF gerado
pelo nosso conversor tem que ter EXATAMENTE as mesmas palavras, na mesma
ordem, que o LibreOffice produz a partir do mesmo .docx (referência de
fidelidade). Nada pode sumir do texto jurídico nem ser inventado. Roda
quando o LibreOffice (soffice) está instalado na máquina de teste."""

import io
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

import docx
import pytest
from pypdf import PdfReader

from app.docx_pdf import docx_para_pdf
from app.documentos_gerados import preencher_contrato_porto_franco, preencher_contrato_rancho_texas, preencher_recibo

LOTE_PF = {"identificador": "LOTE 08 - QUADRA 10", "lote_numero": 8, "quadra": "10", "tamanho_m2": 200.0,
           "valor_total": 89990.0, "entrada": 8999.0, "entrega": 8999.0, "parcela_mensal": 719.92, "qtd_parcelas": 100}
LOTE_RT = {"identificador": "LOTE 012", "lote_numero": 12, "quadra": "12", "tamanho_m2": 810.84,
           "valor_total": 227035.20, "entrada": 22703.52, "entrega": 22703.68, "parcela_mensal": 1816.28, "qtd_parcelas": 100}
CLIENTE = {"nome": "Maria Teste da Silva", "cpf": "01147332444"}
CORRETOR = {"nome": "Corretor Teste", "cpf_cnpj": "12345678909", "creci": "1234", "banco": "Caixa", "agencia": "0001",
            "conta": "12345-6"}


def dq(fp, casado=False):
    d = {
        "proponente": {"nome": "Maria Teste da Silva", "cpf_cnpj": "01147332444", "rg": "123", "orgao_expedidor": "SSP/RN",
                       "nacionalidade": "brasileira", "profissao": "Analista", "data_nascimento": "1981-11-01"},
        "estado_civil": "casado" if casado else "solteiro",
        "endereco_residencial": {"rua": "Rua A", "numero": "10", "bairro": "Centro", "cidade": "Mossoró", "estado": "RN",
                                 "cep": "59600000"},
        "forma_pagamento": fp,
    }
    if casado:
        d["conjuge"] = {"nome": "João Teste", "cpf_cnpj": "52998224725", "rg": "456", "nacionalidade": "brasileiro",
                        "profissao": "Médico", "data_nascimento": "1980-01-01"}
    return d


def fp_unica(lote):
    return {"a_vista": False, "renda": "7.000,00", "valor_proposto": lote["valor_total"],
            "sinal": f"{lote['entrada']:.2f}".replace(".", ","), "sinal_forma": "unica", "sinal_meio": "pix",
            "sinal_vencimento": "2026-10-05", "dividido_em_parcelas": 100,
            "valor_parcela": f"{lote['parcela_mensal']:.2f}".replace(".", ","), "primeiro_mes": "2026-11-10",
            "vencimento": "10", "chave_valor": f"{lote['entrega']:.2f}".replace(".", ","), "confirmado": True}


FP_PARCELADA = {**fp_unica(LOTE_PF), "sinal_forma": "parcelada", "sinal_parcelas": 3, "sinal_meio": "boleto"}
FP_AVISTA = {"a_vista": True, "renda": "10.000,00", "valor_proposto": 89990.0, "avista_meio": "transferencia",
             "avista_data": "2026-10-10", "confirmado": True}
# proposta antiga (antes do formulário de 02/10), igual à 0002 do Lucas
FP_LEGADO = {"renda": "7000", "sinal": "8.999", "a_vista": False, "vencimento": "20", "primeiro_mes": "20/10/2026",
             "valor_parcela": "719,92", "valor_proposto": 89990.0, "dividido_em_parcelas": 100}


def proposta(fp, lote=LOTE_PF, casado=False):
    return {"numero": 99, "numero_contrato": 91, "valor_proposto": fp.get("valor_proposto") or lote["valor_total"],
            "dados_qualificacao": dq(fp, casado)}


def texto_docx(b: bytes) -> str:
    d = docx.Document(io.BytesIO(b))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                partes.append(c.text)
    return re.sub(r"\s+", " ", " ".join(partes))


def texto_pdf(b: bytes) -> str:
    r = PdfReader(io.BytesIO(b))
    return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in r.pages))


def pf(fp, casado=False, **kw):
    return preencher_contrato_porto_franco(proposta=proposta(fp, LOTE_PF, casado), lote=LOTE_PF, cliente=CLIENTE, corretor=CORRETOR,
                                           valor_comissao=4499.5, data_contrato=date(2026, 10, 2), **kw)


def rt(fp, casado=False):
    return preencher_contrato_rancho_texas(proposta=proposta(fp, LOTE_RT, casado), lote=LOTE_RT, cliente=CLIENTE,
                                           data_contrato=date(2026, 10, 2))


# ------------------------------------------------------------------ Porto Franco
def test_pf_entrada_unica():
    t = texto_docx(pf(fp_unica(LOTE_PF)))
    assert ("Entrada / Arras confirmatórias: R$ 8.999,00 (oito mil novecentos e noventa e nove reais), "
            "paga via PIX em 05/10/2026") in t
    assert ("Parcelas mensais: 100 (cem) parcelas mensais e sucessivas de R$ 719,92 (setecentos e dezenove reais e "
            "noventa e dois centavos) cada, vencendo a primeira em 10/11/2026 e as demais no mesmo dia dos meses "
            "subsequentes.") in t
    assert "paga via PIX em 05/10/2026." in t and ", dia 10" not in t
    assert "Chave: R$ 8.999,00 (oito mil novecentos e noventa e nove reais), a ser paga na entrega das chaves." in t
    assert "Maria Teste da Silva" in t and "011.473.324-44" in t


def test_pf_entrada_parcelada():
    t = texto_docx(pf(FP_PARCELADA))
    assert ("paga em 3 (três) parcelas mensais, 2 (duas) de R$ 2.999,67 (dois mil novecentos e noventa e nove "
            "reais e sessenta e sete centavos) e 1 (uma) de R$ 2.999,66 (dois mil novecentos e noventa e nove reais e "
            "sessenta e seis centavos), via boleto bancário, vencendo a primeira em 05/10/2026") in t


def test_pf_a_vista():
    t = texto_docx(pf(FP_AVISTA))
    assert ("Entrada / Arras confirmatórias: R$ 89.990,00 (oitenta e nove mil novecentos e noventa reais), "
            "correspondente ao pagamento à vista do preço, via transferência bancária (TED), em 10/10/2026") in t
    assert "Parcelas mensais: não há (pagamento à vista)." in t
    assert "Chave: não há (pagamento à vista)." in t
    assert "719,92" not in t  # nada do plano de tabela vaza pra venda à vista


def test_pf_legado_continua_funcionando():
    t = texto_docx(pf(FP_LEGADO))
    assert "Entrada / Arras confirmatórias: R$ 8.999,00 (oito mil novecentos e noventa e nove reais)" in t
    assert "vencendo a primeira em 20/10/2026 e as demais no mesmo dia" in t
    assert "R$ 9,00" not in t


def test_pf_nome_na_assinatura_com_mesma_fonte_do_cpf():
    d = docx.Document(io.BytesIO(pf(fp_unica(LOTE_PF))))
    nome, cpf = d.paragraphs[313], d.paragraphs[314]
    assert nome.text == "Maria Teste da Silva"
    assert nome.runs[0]._r.rPr.xml == cpf.runs[0]._r.rPr.xml


# ------------------------------------------------------------------ Rancho Texas
def test_rt_entrada_unica_e_chave():
    t = texto_docx(rt({**fp_unica(LOTE_RT), "chave_vencimento": "2030-09-30"}))
    assert ("Sinal/Arras confirmatórias: R$ 22.703,52 (vinte e dois mil setecentos e três reais e cinquenta e dois "
            "centavos) sendo pago via PIX para o dia 05/10/2026.") in t
    # chave nunca sai com data (09/10), nem se uma data antiga estiver gravada
    assert ("Chave: R$ 22.703,68 (vinte e dois mil setecentos e três reais e sessenta e oito centavos), "
            "a ser paga na entrega das chaves.") in t
    assert "30/09/2030" not in t and "30/09/2029" not in t
    assert "XX" not in t and "08/09/2026" not in t and "20.175,40" not in t
    # medidas de exemplo de outro lote não podem sair no contrato (PDF não se edita)
    for exemplo in ("720,55", "Alameda Georgia", "lote 82", "37,14m", "38,04m"):
        assert exemplo not in t
    assert "__________m de frente com" in t


def test_rt_a_vista():
    fp = {**FP_AVISTA, "valor_proposto": 227035.20}
    t = texto_docx(rt(fp))
    assert "Parcelas: não há (pagamento à vista)." in t
    assert "Chave: não há (pagamento à vista)." in t


def test_rt_casado_tem_compradora():
    t = texto_docx(rt(fp_unica(LOTE_RT), casado=True))
    assert "João Teste" in t and "529.982.247-25" in t


# ------------------------------------------------------------------ PDF
@pytest.mark.parametrize(
    "nome,gerar,trechos",
    [
        ("pf_unica", lambda: pf(fp_unica(LOTE_PF)), ["paga via PIX em 05/10/2026", "CONTRATO Nº 0091 – VERSÃO 1"]),
        ("pf_avista", lambda: pf(FP_AVISTA), ["correspondente ao pagamento à vista do preço"]),
        ("rt_parcelada", lambda: rt({**fp_unica(LOTE_RT), "sinal_forma": "parcelada", "sinal_parcelas": 2}),
         ["sendo pago em 2 (duas) parcelas mensais de R$ 11.351,76"]),
        ("recibo_pf", lambda: preencher_recibo(condominio_nome="Porto Franco Residencial", cliente_nome="Maria Teste",
                                               cliente_cpf="01147332444", valor=8999.0, lote_quadra="10",
                                               lote_numero="8", data=date(2026, 10, 2))[0],
         ["R$ 8.999,00", "011.473.324-44", "02 de Outubro de 2026"]),
        ("recibo_rt", lambda: preencher_recibo(condominio_nome="Rancho Texas", cliente_nome="Maria Teste",
                                               cliente_cpf="01147332444", valor=8999.0, lote_quadra=None,
                                               lote_numero="LOTE 012", data=date(2026, 10, 2))[0],
         ["R$ 8.999,00", "LOTE 012"]),
    ],
)
def test_pdf_tem_o_conteudo(nome, gerar, trechos):
    b = gerar()
    pdf = docx_para_pdf(b)
    assert pdf[:5] == b"%PDF-"
    t = texto_pdf(pdf)
    for trecho in trechos:
        assert trecho in t, (nome, trecho)


def _tokens(texto: str) -> list[str]:
    # hífen U+2010 (LibreOffice) = hífen comum (nosso PDF: a fonte não tem o U+2010)
    return re.findall(r"\w+|[^\w\s]", texto.replace("\u2010", "-"))


def _sem_numeros_de_pagina(tokens: list[str]) -> list[str]:
    return [t for t in tokens if not t.isdigit() or len(t) > 2]


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice não instalado")
@pytest.mark.parametrize(
    "gerar",
    [
        lambda: pf(fp_unica(LOTE_PF)),
        lambda: pf(FP_PARCELADA),
        lambda: rt(fp_unica(LOTE_RT), casado=True),
        lambda: preencher_recibo(condominio_nome="Rancho Texas", cliente_nome="Maria Teste", cliente_cpf="01147332444",
                                 valor=8999.0, lote_quadra=None, lote_numero="LOTE 012", data=date(2026, 10, 2))[0],
        lambda: preencher_recibo(condominio_nome="Porto Franco", cliente_nome="Maria Teste", cliente_cpf="01147332444",
                                 valor=1234567.89, lote_quadra="10", lote_numero="8", data=date(2026, 10, 2))[0],
    ],
)
def test_pdf_tem_exatamente_o_texto_do_libreoffice(gerar):
    b = gerar()
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.docx"
        src.write_bytes(b)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                       check=True, capture_output=True, timeout=180)
        ref = (Path(tmp) / "doc.pdf").read_bytes()
    esperado = _sem_numeros_de_pagina(_tokens(texto_pdf(ref)))
    obtido = _sem_numeros_de_pagina(_tokens(texto_pdf(docx_para_pdf(b))))
    assert obtido == esperado


# ------------------------------------------------------------------ cônjuge / naturalidade (03/10)
def test_rt_uniao_estavel_entra_como_compradora_com_naturalidade():
    d = dq(fp_unica(LOTE_RT), casado=True)
    d["estado_civil"] = "uniao_estavel"
    d["tem_conjuge"] = True
    d["proponente"]["naturalidade"] = "Mossoró/RN"
    d["conjuge"]["naturalidade"] = "Natal/RN"
    prop = {"numero": 99, "numero_contrato": 91, "valor_proposto": LOTE_RT["valor_total"], "dados_qualificacao": d}
    t = texto_docx(preencher_contrato_rancho_texas(proposta=prop, lote=LOTE_RT, cliente=CLIENTE, data_contrato=date(2026, 10, 2)))
    assert "João Teste" in t and "529.982.247-25" in t
    assert "União estável" in t
    assert "LOCAL: Mossoró/RN" in t and "LOCAL: Natal/RN" in t


def test_rt_sem_conjuge_respondido_nao_tem_compradora():
    d = dq(fp_unica(LOTE_RT), casado=True)
    d["tem_conjuge"] = False  # resposta direta manda (dados de cônjuge esquecidos não entram)
    prop = {"numero": 99, "numero_contrato": 91, "valor_proposto": LOTE_RT["valor_total"], "dados_qualificacao": d}
    t = texto_docx(preencher_contrato_rancho_texas(proposta=prop, lote=LOTE_RT, cliente=CLIENTE, data_contrato=date(2026, 10, 2)))
    assert "João Teste" not in t


def test_pf_local_de_nascimento():
    d = dq(fp_unica(LOTE_PF))
    d["proponente"]["naturalidade"] = "Mossoró/RN"
    prop = {"numero": 99, "numero_contrato": 91, "valor_proposto": LOTE_PF["valor_total"], "dados_qualificacao": d}
    t = texto_docx(preencher_contrato_porto_franco(proposta=prop, lote=LOTE_PF, cliente=CLIENTE, corretor=CORRETOR,
                                                   valor_comissao=0, data_contrato=date(2026, 10, 2)))
    assert "LOCAL: Mossoró/RN" in t
    # sem naturalidade: continua em branco, como sempre foi
    assert "LOCAL: Mossoró" not in texto_docx(pf(fp_unica(LOTE_PF)))


# ------------------------------------------------- IX. Proporção / chave sem data
def tabela_proporcao(b: bytes) -> list[list[str]]:
    from app.documentos_gerados import _celulas_unicas, _todas_tabelas
    d = docx.Document(io.BytesIO(b))
    for t in _todas_tabelas(d):
        cab = [c.text.strip() for c in _celulas_unicas(t.rows[0])]
        if len(cab) >= 2 and cab[0].lower().startswith("nome") and "%" in cab[-1]:
            return [[c.text.strip() for c in _celulas_unicas(r)] for r in t.rows[1:]]
    raise AssertionError("tabela IX não encontrada")


def test_pf_proporcao_so_comprador_100():
    linhas = tabela_proporcao(pf(fp_unica(LOTE_PF)))
    assert linhas[0] == ["Maria Teste da Silva", "100%"]
    assert all(c == "" for linha in linhas[1:] for c in linha)


def test_pf_proporcao_com_conjuge_marcado_50_50():
    def com_conjuge(fp):
        p = proposta(fp, LOTE_PF, casado=True)
        p["dados_qualificacao"]["tem_conjuge"] = True
        return preencher_contrato_porto_franco(proposta=p, lote=LOTE_PF, cliente=CLIENTE, corretor=CORRETOR,
                                               valor_comissao=4499.5, data_contrato=date(2026, 10, 2))
    linhas = tabela_proporcao(com_conjuge(fp_unica(LOTE_PF)))
    assert linhas[0] == ["Maria Teste da Silva", "50%"]
    assert linhas[1] == ["João Teste", "50%"]


def test_pf_proporcao_casado_sem_resposta_fica_100():
    # proposta antiga (sem a pergunta "tem cônjuge?") -- só o comprador
    linhas = tabela_proporcao(pf(fp_unica(LOTE_PF), casado=True))
    assert linhas[0] == ["Maria Teste da Silva", "100%"]


def test_pf_chave_sem_data_mesmo_com_data_gravada():
    t = texto_docx(pf({**fp_unica(LOTE_PF), "chave_vencimento": "2026-10-30"}))
    assert "a ser paga na entrega das chaves." in t
    assert "30/10/2026" not in t and "até" not in t.split("Chave:")[1][:120]


def test_rt_parcelas_sem_dia_repetido():
    t = texto_docx(rt(fp_unica(LOTE_RT)))
    assert "vencendo a primeira em 10/11/2026 e as demais todo dia 10 dos meses subsequentes." in t
    assert "CONTRATO Nº 0091 – VERSÃO 1" in t


def test_contrato_sem_numero_nao_sai():
    p = proposta(fp_unica(LOTE_PF))
    p.pop("numero_contrato")
    with pytest.raises(ValueError):
        preencher_contrato_porto_franco(proposta=p, lote=LOTE_PF, cliente=CLIENTE, corretor=CORRETOR,
                                        valor_comissao=0, data_contrato=date(2026, 10, 2))


# ------------------------------------------------- V. Memorial descritivo (09/10)
import json as _json
MEMORIAL_PF = {(m["quadra"], m["lote_numero"]): m for m in _json.load(
    open(Path(__file__).parent.parent / "seed" / "data" / "porto_franco_memorial.json", encoding="utf-8"))}


def test_memorial_tem_os_814_lotes_e_o_lote_6_da_quadra_6_confere():
    assert len(MEMORIAL_PF) == 814
    m = MEMORIAL_PF[("6", 6)]
    assert m["area_texto"] == "Área privativa: 200,00m², Área comum: 90,16m² e Área real: 290,16m²"
    assert m["confrontacoes"] == ["10,00m de frente com a Rua Porto Areia Branca", "10,00m de fundo com o lote 25",
                                  "20,00m de comprimento pelo lado direito com o lote 07",
                                  "20,00m de comprimento pelo lado esquerdo com o lote 05"]


def test_pf_item_v_sai_com_o_memorial():
    lote = {**LOTE_PF, "quadra": "6", "lote_numero": 6, "tamanho_m2": 199.0}
    b = preencher_contrato_porto_franco(proposta=proposta(fp_unica(LOTE_PF)), lote=lote, cliente=CLIENTE,
                                        corretor=CORRETOR, valor_comissao=0, data_contrato=date(2026, 10, 9),
                                        memorial=MEMORIAL_PF[("6", 6)])
    d = docx.Document(io.BytesIO(b))
    ps = [p.text for p in d.tables[0].rows[7].cells[0].paragraphs]
    assert "LOTE N° 6 QUADRA N° 6, com área total de 200,00 m²" in ps[1]  # área = privativa do memorial
    assert ps[2:7] == [MEMORIAL_PF[("6", 6)]["area_texto"], *MEMORIAL_PF[("6", 6)]["confrontacoes"]]
    # mesma fonte do texto do item V, sem negrito
    assert all(not r.bold for p in d.tables[0].rows[7].cells[0].paragraphs[2:7] for r in p.runs)
    t = texto_pdf(docx_para_pdf(b))
    assert "10,00m de fundo com o lote 25" in t and "20,00m de comprimento pelo lado esquerdo com o lote 05" in t


def test_pf_item_v_lote_de_esquina_com_linha_longa():
    m = MEMORIAL_PF[("1", 19)]
    b = preencher_contrato_porto_franco(proposta=proposta(fp_unica(LOTE_PF)),
                                        lote={**LOTE_PF, "quadra": "1", "lote_numero": 19}, cliente=CLIENTE,
                                        corretor=CORRETOR, valor_comissao=0, data_contrato=date(2026, 10, 9),
                                        memorial=m)
    t = texto_docx(b)
    assert "3,65m de frente com a Rua Porto Angra dos Reis + 7,88m com a rotatória da Rua Porto Angra dos Reis" in t


# ------------------------------------------------------------------ Riviera e Veneza (29/10)
from app.documentos_gerados import preencher_contrato_riviera, preencher_contrato_veneza, problemas_modelo_contrato  # noqa: E402

LOTE_RIV = {"identificador": "LOTE 20 - QUADRA A", "lote_numero": 20, "quadra": "A", "tamanho_m2": 200.0,
            "valor_total": 195000.0, "entrada": 9750.0, "entrega": None, "parcela_mensal": 1852.50, "qtd_parcelas": 100}
LOTE_VEN = {"identificador": "CASA 31 - QUADRA 2", "lote_numero": 31, "quadra": "2", "tamanho_m2": 200.0,
            "area_construida_m2": 75.77, "valor_total": 343000.0, "entrada": 17150.0, "entrega": None,
            "parcela_mensal": 1429.17, "qtd_parcelas": 36, "qtd_baloes": 3, "valor_balao": 11433.33,
            "financiamento": 240099.89}


def fp_riviera():
    return {**fp_unica({**LOTE_RIV, "entrega": 0.0}), "chave_valor": None}


def fp_veneza(valor_proposto=343000.0, financiamento="240.099,89"):
    return {"a_vista": False, "renda": "9.000,00", "valor_proposto": valor_proposto, "valor_tabela": 343000.0,
            "sinal": "17.150,00", "sinal_forma": "unica", "sinal_meio": "pix", "sinal_vencimento": "2026-11-05",
            "dividido_em_parcelas": 36, "valor_parcela": "1.429,17", "primeiro_mes": "2026-12-05", "vencimento": "5",
            "balao_qtd": 3, "balao_valor": "11.433,33", "financiamento_valor": financiamento, "confirmado": True}


def riv(fp, casado=False):
    return preencher_contrato_riviera(proposta=proposta(fp, LOTE_RIV, casado), lote=LOTE_RIV, cliente=CLIENTE,
                                      data_contrato=date(2026, 10, 29))


def ven(fp, comissao=17150.0):
    return preencher_contrato_veneza(proposta=proposta(fp, LOTE_VEN), lote=LOTE_VEN, cliente=CLIENTE,
                                     valor_comissao=comissao, data_contrato=date(2026, 10, 29))


def test_riviera_preenche_quadro_e_tira_o_exemplo_de_2015():
    t = texto_docx(riv(fp_riviera()))
    assert "CONTRATO Nº 0091 – VERSÃO 1" in t
    assert "NOME: Maria Teste da Silva" in t and "CPF: 011.473.324-44" in t
    assert "LOTE Nº 20 QUADRA Nº A, com área total de 200,00 m²" in t
    assert "R$ 195.000,00 (cento e noventa e cinco mil reais)" in t
    assert "Entrada: R$ 9.750,00 (nove mil setecentos e cinquenta reais), paga via PIX em 05/10/2026." in t
    assert "Parcelas: 100 (cem) parcelas mensais e sucessivas de R$ 1.852,50" in t
    assert "neste dia 29 de Outubro de 2026." in t
    assert "30/08/2015" not in t and "Parcelas de Obra" not in t and "6.500,00" not in t
    assert "170.000,00" not in t and "22 de Fevereiro" not in t
    d = docx.Document(io.BytesIO(riv(fp_riviera())))
    assert sum(1 for r in d.tables[0].rows if "VII. FORMA DE PAGAMENTO" in r.cells[0].text) == 1
    assert tabela_proporcao(riv(fp_riviera()))[0] == ["Maria Teste da Silva", "100%"]
    i = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("CPF: 011"))
    assert d.paragraphs[i - 1].text == "Maria Teste da Silva" and "____" in d.paragraphs[i - 2].text


def test_riviera_com_conjuge_50_50():
    fp = fp_riviera()
    b = preencher_contrato_riviera(proposta={**proposta(fp, LOTE_RIV, casado=True),
                                             "dados_qualificacao": {**dq(fp, True), "tem_conjuge": True}},
                                   lote=LOTE_RIV, cliente=CLIENTE, data_contrato=date(2026, 10, 29))
    assert tabela_proporcao(b)[:2] == [["Maria Teste da Silva", "50%"], ["João Teste", "50%"]]


def test_veneza_preenche_quadro_resumo():
    t = texto_docx(ven(fp_veneza()))
    assert "CONTRATO Nº 0091 – VERSÃO 1" in t
    assert ("1.2 PROMITENTE COMPRADOR: Maria Teste da Silva, brasileira, Solteiro(a), inscrito no CPF sob nº "
            "011.473.324-44, RG sob nº 123 SSP/RN, com endereço na Rua A, nº 10, Centro, Mossoró/RN, CEP: 59600-000, "
            "doravante denominado simplesmente COMPRADOR.") in t
    assert "Brasileiro, Divorciado" not in t
    assert "casa n.º 31 da quadra 2" in t and "área privativa de 75,77 m²" in t and "fração ideal de ____________" in t
    assert "3.1: DO PREÇO TOTAL DO IMÓVEL: R$ 343.000,00 (trezentos e quarenta e três mil reais)" in t
    assert "A - BÔNUS: R$ 0,00 (zero reais)" in t
    assert "B - ENTRADA: R$ 17.150,00 (dezessete mil cento e cinquenta reais)" in t
    assert "C - PREÇO DO IMÓVEL: R$ 343.000,00" in t
    assert "36 (trinta e seis) parcelas de R$ 1.429,17 (mil quatrocentos e vinte e nove reais e dezessete centavos)" in t
    assert "3 (três) parcelas de R$ 11.433,33 (onze mil quatrocentos e trinta e três reais e trinta e três centavos)" in t
    assert "3.2.1. Financiamento: R$ 240.099,89" in t
    assert "Recurso Próprio: R$ 17.150,00 (dezessete mil cento e cinquenta reais)" in t
    assert "CORRETAGEM: R$ 17.150,00 (dezessete mil cento e cinquenta reais) que corresponde a 5% do valor total" in t
    assert "Mossoró/RN, 29 de Outubro de 2026" in t
    assert "R$ XX" not in t and "XX da quadra" not in t and "(XXX" not in t and "12.500,00" not in t and "299.000,00" not in t and "05 de Dezembro" not in t


def test_veneza_desconto_vira_bonus_e_percentual_da_comissao():
    t = texto_docx(ven(fp_veneza(330000.0, "227.099,89"), comissao=12500.0))
    assert "PREÇO TOTAL DO IMÓVEL: R$ 343.000,00" in t
    assert "A - BÔNUS: R$ 13.000,00 (treze mil reais)" in t
    assert "C - PREÇO DO IMÓVEL: R$ 330.000,00 (trezentos e trinta mil reais)" in t
    assert "que corresponde a 3,64% do valor total" in t


def test_veneza_a_vista():
    fp = {"a_vista": True, "renda": "9.000,00", "valor_proposto": 343000.0, "avista_meio": "pix",
          "avista_data": "2026-11-05", "confirmado": True}
    t = texto_docx(ven(fp))
    assert "D – PARCELAS: não há (pagamento à vista)" in t and "Financiamento: não há (pagamento à vista)" in t


def test_modelo_sem_lugar_para_chave_ou_baloes_trava():
    assert problemas_modelo_contrato("veneza", fp_veneza(), LOTE_VEN) == []
    com_chave = {**fp_veneza(financiamento="230.099,89"), "chave_valor": "10.000,00"}
    assert "chave" in problemas_modelo_contrato("veneza", com_chave, LOTE_VEN)[0]
    assert "balões" in problemas_modelo_contrato("riviera", fp_veneza(), LOTE_VEN)[0]
    assert problemas_modelo_contrato("riviera", fp_riviera(), LOTE_RIV) == []


@pytest.mark.parametrize("gerar", [lambda: riv(fp_riviera()), lambda: ven(fp_veneza())])
def test_riviera_veneza_viram_pdf(gerar):
    pdf = docx_para_pdf(gerar())
    assert pdf[:5] == b"%PDF-"
    t = texto_pdf(pdf)
    assert "Maria Teste da Silva" in t and "011.473.324-44" in t


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice não instalado")
def test_riviera_pdf_tem_exatamente_o_texto_do_libreoffice():
    test_pdf_tem_exatamente_o_texto_do_libreoffice(lambda: riv(fp_riviera()))


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice não instalado")
def test_veneza_pdf_tem_exatamente_o_texto_do_libreoffice():
    """O modelo do Veneza tem um acento a mais em "poderá" (cláusula de
    revalidação) -- erro de digitação do próprio modelo, que não mexemos; cada
    programa extrai esse caractere de um jeito, então essa palavra fica de fora
    da comparação."""
    import unicodedata as ud
    b = ven(fp_veneza())
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.docx"
        src.write_bytes(b)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                       check=True, capture_output=True, timeout=180)
        ref = (Path(tmp) / "doc.pdf").read_bytes()

    def limpa(tokens):
        return [t for t in tokens if not t.startswith("poder") and not all(ud.combining(ch) for ch in t)]

    esperado = limpa(_sem_numeros_de_pagina(_tokens(texto_pdf(ref))))
    obtido = limpa(_sem_numeros_de_pagina(_tokens(texto_pdf(docx_para_pdf(b)))))
    assert obtido == esperado


# ------------------------------------------------------------------ nome que sumia do contrato (29/10)
NOME_ESTILIZADO = "\U0001d40c\U0001d41a\U0001d42b\U0001d422\U0001d41a ​Teste  da Silva "  # 𝐌𝐚𝐫𝐢𝐚 + zero-width + NBSP


def test_limpar_nome():
    from app.texto_seguro import limpar_nome
    assert limpar_nome(NOME_ESTILIZADO) == "Maria Teste da Silva"
    assert limpar_nome("﻿​  ") == ""
    assert limpar_nome("José d'Ávila") == "José d'Ávila"


@pytest.mark.parametrize("gerar_nome", ["pf", "rt", "riv", "ven"])
def test_nome_estilizado_sai_no_contrato_e_na_assinatura(gerar_nome):
    fp = {"pf": fp_unica(LOTE_PF), "rt": fp_unica(LOTE_RT), "riv": fp_riviera(), "ven": fp_veneza()}[gerar_nome]
    lote = {"pf": LOTE_PF, "rt": LOTE_RT, "riv": LOTE_RIV, "ven": LOTE_VEN}[gerar_nome]
    p = proposta(fp, lote)
    p["dados_qualificacao"]["proponente"]["nome"] = NOME_ESTILIZADO
    kw = dict(proposta=p, lote=lote, cliente=CLIENTE, data_contrato=date(2026, 10, 29))
    b = {"pf": lambda: preencher_contrato_porto_franco(corretor=CORRETOR, valor_comissao=0, **kw),
         "rt": lambda: preencher_contrato_rancho_texas(**kw),
         "riv": lambda: preencher_contrato_riviera(**kw),
         "ven": lambda: preencher_contrato_veneza(valor_comissao=0, **kw)}[gerar_nome]()
    t = texto_pdf(docx_para_pdf(b))
    assert t.count("Maria Teste da Silva") >= 2, gerar_nome  # quadro resumo + assinatura


def test_nome_so_com_espacos_invisiveis_usa_o_do_cadastro():
    p = proposta(fp_unica(LOTE_PF))
    p["dados_qualificacao"]["proponente"]["nome"] = "​  "
    b = preencher_contrato_porto_franco(proposta=p, lote=LOTE_PF, cliente={"nome": "Cliente Do Cadastro", "cpf": "01147332444"},
                                        corretor=CORRETOR, valor_comissao=0, data_contrato=date(2026, 10, 2))
    assert texto_docx(b).count("Cliente Do Cadastro") >= 2


def test_pdf_nao_apaga_letra_estilizada_mesmo_sem_limpeza():
    from app.texto_seguro import ajustar_para_cobertura, cobertura_padrao
    assert ajustar_para_cobertura(NOME_ESTILIZADO, cobertura_padrao()).strip().startswith("Maria")
    assert ajustar_para_cobertura("Nº 1, 200,00 m² – “ok”", cobertura_padrao()) == "Nº 1, 200,00 m² – “ok”"
